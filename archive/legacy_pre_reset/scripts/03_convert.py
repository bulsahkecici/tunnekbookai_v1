from __future__ import annotations

import csv
import json
import os
import re
import shutil
import time
import uuid
import zipfile
from pathlib import Path

from utils import (
    find_libreoffice, iso_time, load_config, native_path, read_csv, run_checked, setup_logging, sha256_file,
    write_csv,
)


MANIFEST_FIELDS = [
    "document_id", "source_file", "source_relative_path", "source_extension", "source_sha256",
    "output_md", "converter", "status", "started_at", "finished_at", "duration_seconds",
    "error_message", "duplicate_status", "preferred_variant",
]
MODERN_EXTENSIONS = {".pdf", ".docx", ".pptx", ".pptm", ".xlsx", ".txt", ".jpg", ".jpeg", ".png"}
LEGACY_MAP = {".doc": ".docx", ".ppt": ".pptx", ".xls": ".xlsx", ".rtf": ".docx"}
_DOCLING_CONVERTER = None
_OCR_ENGINE = None


def as_bool(value) -> bool:
    return str(value).casefold() in {"true", "1", "yes"}


def conversion_maps(root: Path):
    exact_rows = read_csv(root / "data/duplicates/exact_duplicates.csv")
    logical_rows = read_csv(root / "data/duplicates/logical_duplicates.csv")
    exact = {
        row["document_id"]: (row["duplicate_group_id"], as_bool(row["canonical_candidate"]))
        for row in exact_rows
    }
    logical = {
        row["document_id"]: (row["logical_group_id"], as_bool(row["preferred_source"]), row.get("confidence", ""))
        for row in logical_rows
    }
    return exact, logical


def markdown_path(root: Path, relative_path: str) -> Path:
    return (root / "data/markdown" / Path(relative_path)).with_suffix(".md")


def docling_convert(source: Path) -> str:
    global _DOCLING_CONVERTER
    cache_root = load_config()["project_root"] / "data/temp/model_cache"
    cache_root.mkdir(parents=True, exist_ok=True)
    os.environ.setdefault("HF_HOME", str(cache_root / "huggingface"))
    os.environ.setdefault("TORCH_HOME", str(cache_root / "torch"))
    # Windows hosts without Visual Studio C++ cannot use TorchInductor's generated C++.
    # Eager mode is slower but reliable and still uses Docling's actual models.
    os.environ.setdefault("TORCHDYNAMO_DISABLE", "1")
    try:
        import torch  # type: ignore
        torch._dynamo.config.suppress_errors = True
    except Exception:
        pass
    from docling.document_converter import DocumentConverter  # type: ignore
    if _DOCLING_CONVERTER is None:
        _DOCLING_CONVERTER = DocumentConverter()
    result = _DOCLING_CONVERTER.convert(str(source))
    return result.document.export_to_markdown()


def fallback_pdf(source: Path) -> str:
    from pypdf import PdfReader
    reader = PdfReader(str(source), strict=False)
    parts = [f"# {source.stem}"]
    for number, page in enumerate(reader.pages, 1):
        text = (page.extract_text() or "").strip()
        if text:
            parts.extend(["", f"## Sayfa {number}", "", text])
    return "\n".join(parts)


def fallback_docx(source: Path) -> str:
    from docx import Document
    document = Document(str(source))
    output = [f"# {source.stem}"]
    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        if not text:
            continue
        style = (paragraph.style.name or "").casefold() if paragraph.style else ""
        match = re.search(r"heading\s*(\d+)", style)
        output.append(("#" * min(6, int(match.group(1))) + " " if match else "") + text)
    for table in document.tables:
        rows = [[cell.text.replace("\n", " ").strip() for cell in row.cells] for row in table.rows]
        if not rows:
            continue
        output.extend(["", "| " + " | ".join(rows[0]) + " |", "| " + " | ".join("---" for _ in rows[0]) + " |"])
        output.extend("| " + " | ".join(row) + " |" for row in rows[1:])
    return "\n\n".join(output)


def fallback_pptx(source: Path) -> str:
    from pptx import Presentation
    presentation = Presentation(str(source))
    output = [f"# {source.stem}"]
    for number, slide in enumerate(presentation.slides, 1):
        texts = []
        for shape in slide.shapes:
            if hasattr(shape, "text") and shape.text.strip():
                texts.append(shape.text.strip())
        output.extend(["", f"## Slayt {number}", ""])
        output.extend(texts)
    return "\n".join(output)


def fallback_xlsx(source: Path) -> str:
    from openpyxl import load_workbook
    workbook = load_workbook(source, read_only=True, data_only=True)
    output = [f"# {source.stem}"]
    for sheet in workbook.worksheets:
        output.extend(["", f"## {sheet.title}", ""])
        rows = []
        for values in sheet.iter_rows(values_only=True):
            rendered = ["" if value is None else str(value).replace("|", "\\|").replace("\n", " ") for value in values]
            if any(rendered):
                rows.append(rendered)
        if rows:
            width = max(len(row) for row in rows)
            rows = [row + [""] * (width - len(row)) for row in rows]
            output.append("| " + " | ".join(rows[0]) + " |")
            output.append("| " + " | ".join("---" for _ in range(width)) + " |")
            output.extend("| " + " | ".join(row) + " |" for row in rows[1:])
    workbook.close()
    return "\n".join(output)


def get_ocr_engine():
    global _OCR_ENGINE
    from rapidocr import RapidOCR
    from rapidocr.utils.typings import EngineType
    if _OCR_ENGINE is None:
        params = {
            "Det.engine_type": EngineType.TORCH,
            "Cls.engine_type": EngineType.TORCH,
            "Rec.engine_type": EngineType.TORCH,
        }
        _OCR_ENGINE = RapidOCR(params=params)
    return _OCR_ENGINE


def ocr_text(image) -> list[str]:
    result = get_ocr_engine()(image)
    return [text.strip() for text in (getattr(result, "txts", None) or []) if text and text.strip()]


def image_convert(source: Path) -> tuple[str, str]:
    """Conservative image caption plus local OCR; no LLM or external API is used."""
    from PIL import Image
    with Image.open(source) as image:
        width, height = image.size
    texts = ocr_text(source)
    clean_name = re.sub(r"[_-]+", " ", source.stem).strip()
    caption = f"Kaynak arşiv görseli: {clean_name}."
    parts = [
        f"# {source.stem}", "", f"![{clean_name}]({source.name})", "",
        "## Konservatif görsel açıklaması", "", caption, "",
        f"- Boyut: {width} × {height} piksel",
        "- Açıklama yöntemi: dosya adı + yerel OCR; görsel anlamı tahmin edilmedi.",
        "- İnsan görsel incelemesi: gerekli", "", "## OCR metni", "",
    ]
    parts.append("\n\n".join(texts) if texts else "OCR ile metin tespit edilmedi.")
    return "\n".join(parts), "rapidocr_torch+filename_caption"


def pdf_rapidocr(source: Path, root: Path, logger) -> str:
    """Render every page locally and OCR it; intended for PDFs with no usable text layer."""
    import numpy as np
    import pypdfium2 as pdfium
    scale = 1.2
    cache_dir = root / "data/temp/ocr_page_cache"
    cache_dir.mkdir(parents=True, exist_ok=True)
    cache_path = cache_dir / f"{sha256_file(source)}-scale12.jsonl"
    cached: dict[int, list[str]] = {}
    if cache_path.exists():
        for line in cache_path.read_text(encoding="utf-8", errors="replace").splitlines():
            try:
                item = json.loads(line)
                cached[int(item["page"])] = list(item.get("texts") or [])
            except Exception:
                continue
    document = pdfium.PdfDocument(str(source))
    output = [f"# {source.stem}"]
    try:
        for index in range(len(document)):
            page_number = index + 1
            texts = cached.get(page_number)
            if texts is None:
                page = document[index]
                bitmap = page.render(scale=scale, rotation=0)
                image = bitmap.to_pil()
                texts = ocr_text(np.asarray(image))
                page.close()
                with cache_path.open("a", encoding="utf-8") as handle:
                    handle.write(json.dumps({"page": page_number, "texts": texts}, ensure_ascii=False) + "\n")
            if texts:
                output.extend(["", f"## Sayfa {page_number}", "", "\n\n".join(texts)])
            if page_number % 10 == 0:
                logger.info("RapidOCR page checkpoint %d/%d: %s", page_number, len(document), source.name)
    finally:
        document.close()
    return "\n".join(output)


def native_convert(source: Path, extension: str) -> tuple[str, str]:
    if extension == ".pdf":
        return fallback_pdf(source), "pypdf_fallback"
    if extension == ".docx":
        return fallback_docx(source), "python_docx_fallback"
    if extension in {".pptx", ".pptm"}:
        return fallback_pptx(source), "python_pptx_fallback"
    if extension == ".xlsx":
        return fallback_xlsx(source), "openpyxl_fallback"
    if extension == ".txt":
        return "# " + source.stem + "\n\n" + source.read_text(encoding="utf-8-sig", errors="replace"), "text_builtin"
    if extension in {".jpg", ".jpeg", ".png"}:
        return image_convert(source)
    raise ValueError(f"No native fallback for {extension}")


def docling_eligible(source: Path, extension: str) -> bool:
    """Bound CPU-heavy Docling work while preserving complete native conversion."""
    size_mb = source.stat().st_size / 1048576
    if extension == ".pdf":
        try:
            from pypdf import PdfReader
            return size_mb <= 5 and len(PdfReader(str(source), strict=False).pages) <= 3
        except Exception:
            return size_mb <= 1
    if extension == ".docx":
        return size_mb <= 1
    if extension in {".pptx", ".pptm"}:
        return size_mb <= 2
    if extension == ".xlsx":
        return size_mb <= 1
    return False


def convert_one(source: Path, extension: str, logger) -> tuple[str, str]:
    if extension in {".jpg", ".jpeg", ".png"}:
        return image_convert(source)
    if docling_eligible(source, extension):
        try:
            text = docling_convert(source)
            return text, "docling"
        except ImportError:
            logger.debug("Docling not installed; native fallback for %s", source)
        except Exception as exc:
            logger.warning("Docling failed for %s; trying native fallback: %s", source, exc)
    return native_convert(source, extension)


OFFICE_FILTERS = {
    ".docx": "docx:Office Open XML Text",
    ".pptx": "pptx:Impress MS PowerPoint 2007 XML",
    ".xlsx": "xlsx:Calc MS Excel 2007 XML",
}


def detected_ooxml_extension(source: Path) -> str | None:
    try:
        if source.read_bytes()[:2] != b"PK":
            return None
        with zipfile.ZipFile(source) as archive:
            names = set(archive.namelist())
        if any(name.startswith("word/") for name in names):
            return ".docx"
        if any(name.startswith("ppt/") for name in names):
            return ".pptx"
        if any(name.startswith("xl/") for name in names):
            return ".xlsx"
    except Exception:
        return None
    return None


def libreoffice_run(libreoffice: Path, root: Path, args: list[str], timeout: int = 180) -> object:
    profile = root / "data/temp/libreoffice_profiles" / uuid.uuid4().hex
    profile.mkdir(parents=True, exist_ok=True)
    env = os.environ.copy()
    env["SAL_DISABLE_OPENCL"] = "1"
    env["SAL_USE_VCLPLUGIN"] = "svp"
    env["PATH"] = str(libreoffice.parent) + os.pathsep + env.get("PATH", "")
    command = [
        str(libreoffice), f"-env:UserInstallation={profile.resolve().as_uri()}",
        "--headless", "--nologo", "--nodefault", "--nolockcheck", "--norestore", *args,
    ]
    try:
        return run_checked(command, timeout=timeout, env=env)
    finally:
        shutil.rmtree(profile, ignore_errors=True)


def convert_legacy(source: Path, relative_path: str, extension: str, root: Path, libreoffice: Path) -> tuple[Path | None, str, str]:
    target_ext = LEGACY_MAP[extension]
    output_dir = root / "data/converted_office" / Path(relative_path).parent
    output_dir.mkdir(parents=True, exist_ok=True)
    expected = output_dir / (source.stem + target_ext)
    if detected_ooxml_extension(source) == target_ext:
        shutil.copy2(source, expected)
        return expected, "", "ooxml_magic_copy"
    result = libreoffice_run(
        libreoffice, root,
        ["--convert-to", OFFICE_FILTERS[target_ext], "--outdir", str(output_dir), str(source)],
    )
    if result.returncode != 0 or not expected.exists():
        message = (result.stderr or result.stdout or "LibreOffice did not produce the expected output").strip()
        return None, message, "libreoffice"
    return expected, "", "libreoffice"


def meaningful_text(text: str, title: str) -> bool:
    stripped = text.strip()
    if not stripped:
        return False
    without_title = re.sub(rf"^#\s*{re.escape(title)}\s*", "", stripped, flags=re.IGNORECASE).strip()
    return len(re.sub(r"\W+", "", without_title, flags=re.UNICODE)) >= 20


def libreoffice_to_pdf(source: Path, root: Path) -> tuple[Path | None, str]:
    libreoffice = find_libreoffice()
    if not libreoffice:
        return None, "LibreOffice soffice executable was not found"
    output_dir = root / "data/temp/ocr_recovery" / source.parent.name
    output_dir.mkdir(parents=True, exist_ok=True)
    expected = output_dir / (source.stem + ".pdf")
    result = libreoffice_run(
        libreoffice, root,
        ["--convert-to", "pdf:writer_pdf_Export", "--outdir", str(output_dir), str(source)],
    )
    if result.returncode != 0 or not expected.exists():
        return None, (result.stderr or result.stdout or "LibreOffice PDF recovery output missing").strip()
    return expected, ""


def recover_empty_output(source: Path, extension: str, root: Path, logger) -> tuple[str, str, str]:
    """Force layout/OCR for files whose fast converter produced no meaningful text."""
    try:
        recovery_source = source
        prefix = ""
        if extension == ".docx":
            recovery_source, error = libreoffice_to_pdf(source, root)
            if not recovery_source:
                return "", "libreoffice_pdf_recovery", error
            prefix = "libreoffice_pdf+"
        if extension in {".pdf", ".docx"}:
            text = pdf_rapidocr(recovery_source, root, logger)
            return text, prefix + "pypdfium2+rapidocr_torch", ""
    except Exception as exc:
        logger.exception("Forced OCR recovery failed for %s", source)
        return "", "docling_forced_ocr", f"{type(exc).__name__}: {exc}"
    return "", "", "No recovery strategy configured"


def previous_successes(manifest_path: Path) -> dict[str, dict[str, str]]:
    return {
        row["document_id"]: row for row in read_csv(manifest_path)
        if row.get("status") == "success"
    }


def main() -> int:
    config = load_config()
    logger = setup_logging("03_convert", config)
    root, source_root = config["project_root"], config["source_root"]
    inventory = read_csv(root / "data/inventory/inventory.csv")
    if not inventory:
        logger.error("Inventory is missing or empty")
        return 2
    exact, logical = conversion_maps(root)
    manifest_path = root / "data/metadata/conversion_manifest.csv"
    previous = previous_successes(manifest_path)
    manifest: list[dict] = []
    libreoffice = find_libreoffice()
    logger.info("Conversion started: %d inventory rows; Docling and native fallbacks enabled; LibreOffice=%s", len(inventory), libreoffice or "not found")
    for row in inventory:
        started_clock, started_at = time.monotonic(), iso_time()
        extension = row["extension"].lower()
        source = Path(row["absolute_path"])
        output = markdown_path(root, row["relative_path"])
        record = {
            "document_id": row["document_id"], "source_file": row["file_name"],
            "source_relative_path": row["relative_path"], "source_extension": extension,
            "source_sha256": row["sha256"], "output_md": str(output), "converter": "",
            "status": "", "started_at": started_at, "finished_at": "", "duration_seconds": 0,
            "error_message": "", "duplicate_status": "unique", "preferred_variant": True,
        }
        try:
            if row["document_id"] in exact and not exact[row["document_id"]][1]:
                record.update(status="skipped_exact_duplicate", duplicate_status=exact[row["document_id"]][0], preferred_variant=False)
            elif row["document_id"] in logical and not logical[row["document_id"]][1] and logical[row["document_id"]][2] == "high":
                record.update(status="skipped_preferred_variant", duplicate_status=logical[row["document_id"]][0], preferred_variant=False)
            elif row["file_name"].startswith("~$"):
                record.update(status="unsupported", error_message="Temporary Microsoft Office lock file; not a document")
            elif row["document_id"] in previous and previous[row["document_id"]].get("source_sha256") == row["sha256"] and native_path(output).exists():
                record = previous[row["document_id"]].copy()
                manifest.append(record)
                continue
            elif extension in LEGACY_MAP:
                if not libreoffice:
                    record.update(status="requires_libreoffice", converter="libreoffice", error_message="LibreOffice soffice executable was not found")
                else:
                    modern, error, modernization_engine = convert_legacy(source, row["relative_path"], extension, root, libreoffice)
                    if not modern:
                        record.update(status="failed", converter="libreoffice", error_message=error)
                    else:
                        text, engine = convert_one(modern, LEGACY_MAP[extension], logger)
                        if not meaningful_text(text, modern.stem):
                            recovered, recovery_engine, recovery_error = recover_empty_output(modern, LEGACY_MAP[extension], root, logger)
                            if meaningful_text(recovered, modern.stem):
                                text, engine = recovered, recovery_engine
                            else:
                                record.update(status="empty_output", converter=f"{modernization_engine}+{engine}", error_message=recovery_error or "Converter produced no meaningful text")
                        if not record["status"]:
                            native_path(output.parent).mkdir(parents=True, exist_ok=True)
                            native_path(output).write_text(text.strip() + "\n", encoding="utf-8")
                            record.update(status="success", converter=f"{modernization_engine}+{engine}")
            elif extension in MODERN_EXTENSIONS:
                text, engine = convert_one(source, extension, logger)
                if not meaningful_text(text, source.stem):
                    recovered, recovery_engine, recovery_error = recover_empty_output(source, extension, root, logger)
                    if meaningful_text(recovered, source.stem):
                        text, engine = recovered, recovery_engine
                    else:
                        record.update(status="empty_output", converter=engine, error_message=recovery_error or "Converter produced no meaningful text")
                if not record["status"]:
                    native_path(output.parent).mkdir(parents=True, exist_ok=True)
                    native_path(output).write_text(text.strip() + "\n", encoding="utf-8")
                    record.update(status="success", converter=engine)
            else:
                record.update(status="unsupported", error_message=f"No converter configured for {extension or '[no extension]'}")
        except Exception as exc:
            record.update(status="failed", error_message=f"{type(exc).__name__}: {exc}")
            logger.exception("Conversion failed for %s", source)
        record["finished_at"] = iso_time()
        record["duration_seconds"] = round(time.monotonic() - started_clock, 3)
        manifest.append(record)
        if len(manifest) % 10 == 0:
            write_csv(manifest_path, manifest, MANIFEST_FIELDS)
            logger.info("Checkpoint: %d/%d", len(manifest), len(inventory))
    write_csv(manifest_path, manifest, MANIFEST_FIELDS)
    logger.info("Conversion completed: %d records", len(manifest))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
