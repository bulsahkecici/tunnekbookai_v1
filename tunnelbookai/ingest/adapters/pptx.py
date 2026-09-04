"""PPTX adapter (task §13, §14).

Slide order, titles, body text, tables, figures, chart metadata and speaker notes come from
python-pptx, which reads the original package deterministically. Docling runs alongside it
to supply the normalized Markdown and its own table/picture items.

Slide snapshots go through the OfficeRenderer (PPTX -> temp PDF -> slides/slide_NNNN.png).
A missing renderer is non-blocking (§14).
"""

from __future__ import annotations

import zipfile
from pathlib import Path
from typing import Any

from .. import docling_adapter
from ..assets import base_asset, sha256_path
from ..assets import charts as charts_assets
from ..assets import figures as figures_assets
from ..assets import tables as tables_assets
from ..elements import ElementBuilder, FIGURE_REF, HEADING, PARAGRAPH, TABLE_REF
from ..extraction import ExtractionResult, write_normalized
from ..ocr import provider as ocr_provider
from ..office_renderer import VISUAL_RENDERER_UNAVAILABLE
from ..paths import relpath
from .docx import RASTER_EXTENSIONS
from . import AdapterContext

SLIDES_DIRNAME = "slides"


def _shape_text(shape: Any) -> str:
    frame = getattr(shape, "text_frame", None)
    if frame is None:
        return ""
    return "\n".join(p.text for p in frame.paragraphs).strip()


def _chart_metadata(shape: Any) -> dict[str, Any]:
    chart = getattr(shape, "chart", None)
    if chart is None:
        return {}
    meta: dict[str, Any] = {"chart_type": str(getattr(chart, "chart_type", None))}
    try:
        meta["title"] = chart.chart_title.text_frame.text if chart.has_title else None
    except Exception:
        meta["title"] = None
    try:
        meta["categories"] = [str(c) for c in list(chart.plots[0].categories)]
    except Exception:
        meta["categories"] = []
    try:
        meta["series"] = [{"name": s.name, "values": list(s.values)} for s in chart.series]
    except Exception:
        meta["series"] = []
    return meta


def _extract_slide_media(shape: Any, out_dir: Path, index: int) -> tuple[Path | None, str | None]:
    image = getattr(shape, "image", None)
    if image is None:
        return None, None
    suffix = f".{(image.ext or 'png').lower()}"
    if suffix not in RASTER_EXTENSIONS:
        return None, suffix
    out_dir.mkdir(parents=True, exist_ok=True)
    target = out_dir / f"FIG{index:04d}{suffix}"
    target.write_bytes(image.blob)
    return target, suffix


def _parse_pptx(source: Path, bundle: Path, document_id: str,
                result: ExtractionResult) -> tuple[list[dict[str, Any]], list[dict[str, Any]],
                                                   list[dict[str, Any]], list[dict[str, Any]],
                                                   list[dict[str, Any]]]:
    """Returns (slides, figures, tables, charts, elements)."""
    try:
        from pptx import Presentation
    except ImportError:
        result.warn("PPTX_NATIVE_PARSER_UNAVAILABLE")
        return [], [], [], [], []

    try:
        presentation = Presentation(str(source))
    except Exception as exc:
        result.warn(f"PPTX_NATIVE_PARSE_FAILED:{type(exc).__name__}")
        return [], [], [], [], []

    builder = ElementBuilder()
    slides: list[dict[str, Any]] = []
    figures: list[dict[str, Any]] = []
    tables: list[dict[str, Any]] = []
    charts: list[dict[str, Any]] = []
    figure_index = table_index = chart_index = 0
    figures_dir = bundle / figures_assets.FIGURES_DIRNAME

    for slide_number, slide in enumerate(presentation.slides, start=1):
        title = None
        try:
            if slide.shapes.title is not None:
                title = (slide.shapes.title.text or "").strip() or None
        except Exception:
            title = None

        body_parts: list[str] = []
        slide_tables: list[str] = []
        slide_figures: list[str] = []
        slide_charts: list[str] = []

        builder.add(HEADING, title or f"Slide {slide_number}", level=1,
                    slide_number=slide_number)

        for shape in slide.shapes:
            try:
                if getattr(shape, "has_chart", False):
                    chart_index += 1
                    record = charts_assets.record_pptx_chart(
                        document_id, chart_index, bundle, slide_number=slide_number,
                        metadata=_chart_metadata(shape))
                    charts.append(record)
                    slide_charts.append(record["chart_id"])
                    continue
                if getattr(shape, "has_table", False):
                    table_index += 1
                    tid = f"TABLE{table_index:04d}"
                    grid = [[cell.text for cell in row.cells] for row in shape.table.rows]
                    payload = {
                        "document_id": document_id, "table_id": tid, "slide_number": slide_number,
                        "page": None, "caption": None, "bbox": None,
                        "rows": len(grid), "columns": len(grid[0]) if grid else 0,
                        "rectangular": tables_assets.is_rectangular(grid), "grid": grid,
                        "otsl": None, "html": None, "markdown": None, "element_ref": None,
                    }
                    json_path, csv_path = tables_assets.write_table_files(
                        bundle / tables_assets.TABLES_DIRNAME, tid, grid, payload, emit_csv=True)
                    record = {
                        "document_id": document_id, "asset_id": tid, "asset_type": "TABLE",
                        "table_id": tid, "slide_number": slide_number, "page": None,
                        "caption": None, "bbox": None,
                        "rows": payload["rows"], "columns": payload["columns"],
                        "rectangular": payload["rectangular"],
                        "path": relpath(json_path), "structured_path": relpath(json_path),
                        "csv_path": relpath(csv_path) if csv_path else None,
                        "image_path": None, "element_ref": None, "markdown": None,
                        "sha256": sha256_path(json_path),
                    }
                    tables.append(record)
                    slide_tables.append(tid)
                    builder.add(TABLE_REF, "", slide_number=slide_number, asset_id=tid)
                    continue
                if shape.shape_type is not None and getattr(shape, "image", None) is not None:
                    candidate_index = figure_index + 1
                    path, suffix = _extract_slide_media(shape, figures_dir, candidate_index)
                    if path is None:
                        if suffix:
                            result.warn(f"PPTX_VECTOR_MEDIA_SKIPPED:{suffix}")
                        continue
                    figure_index = candidate_index
                    record = figures_assets.register_embedded_image(
                        document_id, figure_index, path, slide_number=slide_number,
                        source_order=figure_index, heading_path=builder.heading_path)
                    record["origin"] = "pptx_shape"
                    figures.append(record)
                    slide_figures.append(record["asset_id"])
                    builder.add(FIGURE_REF, "", slide_number=slide_number,
                                asset_id=record["asset_id"])
                    continue
            except Exception as exc:
                result.warn(f"PPTX_SHAPE_FAILED:{slide_number}:{type(exc).__name__}")
                continue

            text = _shape_text(shape)
            if text and text != title:
                body_parts.append(text)
                builder.add(PARAGRAPH, text, slide_number=slide_number)

        notes = None
        try:
            if slide.has_notes_slide:
                notes = (slide.notes_slide.notes_text_frame.text or "").strip() or None
        except Exception:
            notes = None

        slides.append({
            "slide_number": slide_number,
            "title": title,
            "text": "\n\n".join(body_parts) or None,
            "tables": slide_tables,
            "figures": slide_figures,
            "charts": slide_charts,
            "notes": notes,
        })

    return slides, figures, tables, charts, builder.elements


def run(context: AdapterContext) -> ExtractionResult:
    fmt = context.detection.fmt.value
    result = ExtractionResult(document_id=context.document_id, format=fmt, adapter="pptx")
    source = context.original_path

    slides, figures, tables, charts, elements = _parse_pptx(
        source, context.bundle, context.document_id, result)

    conversion = docling_adapter.convert(source, context.config, input_format=fmt,
                                         do_ocr=False, want_vision=False)
    result.engine = {
        "docling_status": conversion.status,
        "docling_versions": conversion.versions,
        "native_parser": "python-pptx",
        "structural_authority": "original_pptx",
    }
    for error in conversion.errors:
        result.warn(f"DOCLING:{error[:120]}")
    if not conversion.ok and not slides:
        result.fail(f"PPTX_EXTRACTION_FAILED:{conversion.status}")
        return result
    if not conversion.ok:
        result.warn(f"DOCLING_CONVERSION_FAILED:{conversion.status}")

    snapshots: list[dict[str, Any]] = []
    renderer = context.office_renderer
    if context.snapshots_enabled and renderer is not None:
        snapshots, snapshot_warnings = renderer.render_pages(
            source, context.bundle, document_id=context.document_id, scale=context.snapshot_scale,
            dirname=SLIDES_DIRNAME, prefix="slide", kind="SLIDE", id_prefix="SLIDE")
        for warning in snapshot_warnings:
            result.warn(warning)
    elif context.snapshots_enabled:
        result.warn(VISUAL_RENDERER_UNAVAILABLE)
    snapshot_by_slide = {s["page_number"]: s for s in snapshots}
    for slide in slides:
        snapshot = snapshot_by_slide.get(slide["slide_number"])
        slide["snapshot_asset_id"] = snapshot["asset_id"] if snapshot else None
        slide["snapshot_path"] = snapshot["path"] if snapshot else None

    if context.ocr is not None and context.do_ocr and figures:
        for warning in ocr_provider.ocr_figures(context.ocr, figures, context.root):
            result.warn(warning)

    lines: list[str] = []
    for slide in slides:
        lines.append(f"## Slayt {slide['slide_number']}: {slide['title'] or ''}".rstrip(": "))
        if slide["text"]:
            lines += ["", slide["text"]]
        if slide["notes"]:
            lines += ["", f"> Konuşmacı notu: {slide['notes']}"]
        lines.append("")
    native_markdown = "\n".join(lines).strip()
    markdown = conversion.markdown.strip() or native_markdown
    text = "\n\n".join(
        part for slide in slides
        for part in (slide["title"], slide["text"], slide["notes"]) if part
    ).strip() or (conversion.text or "")

    document_json = {
        "document_id": context.document_id,
        "format": fmt,
        "adapter": "pptx",
        "source_filename": context.original_filename,
        "structural_authority": "original_pptx",
        "slide_count": len(slides),
        "slides": slides,
        "elements": elements,
        "docling": conversion.document_dict or None,
    }
    write_normalized(context.bundle, result, markdown=markdown,
                     document_json=document_json, text=text)

    result.slides = slides
    result.pages = snapshots
    result.figures = figures
    result.tables = tables
    result.charts = charts
    result.text_elements = elements
    result.page_count = len(slides)
    result.capabilities = {
        "text": bool(text.strip()),
        "tables": bool(tables),
        "figures": bool(figures),
        "page_snapshots": bool(snapshots),
        "ocr": any(f.get("ocr_status") == "SUCCESS" for f in figures),
        "vision": any(f.get("visual_description_status") == "SUCCESS" for f in figures),
        "formulas": False,
        "charts": bool(charts),
        "slides": bool(slides),
        "sheets": False,
    }
    return result
