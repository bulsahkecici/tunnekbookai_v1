"""XLSX adapter (task §15, §16, §17, §18).

A workbook is NOT a Markdown document. openpyxl reads it twice:

    data_only=False -> formulas preserved verbatim  (the authority)
    data_only=True  -> cached values, when Excel stored them

If no cached value exists the value stays `null`. A formula's own string is never presented
as its value, and a value is never invented (§16, stop-condition §90 "Excel formulas silently
lost").

Per-sheet output (§17):

    sheets/<Sheet>.json   authoritative — cells, formulas, merges, dimensions
    sheets/<Sheet>.csv    only when the sheet is safely rectangular values

Visual sheet snapshots are optional and non-authoritative (§18); a rendering failure never
damages a successful structural extraction.
"""

from __future__ import annotations

import csv
import io
import json
import re
from pathlib import Path
from typing import Any

from ..assets import asset_id, base_asset
from ..elements import ElementBuilder, HEADING, PARAGRAPH
from ..extraction import ExtractionResult, write_normalized
from ..office_renderer import VISUAL_RENDERER_UNAVAILABLE
from ..paths import relpath
from . import AdapterContext

SHEETS_DIRNAME = "sheets"
SHEET_SNAPSHOT_DIRNAME = "sheet_snapshots"
MAX_CELLS_PER_SHEET = 200_000
_SAFE_NAME = re.compile(r"[^0-9A-Za-zÇĞİÖŞÜçğıöşü _.-]+")


def safe_sheet_filename(name: str, index: int) -> str:
    cleaned = _SAFE_NAME.sub("_", name).strip() or f"Sheet{index}"
    return cleaned[:80]


def _cell_record(cell: Any, cached: Any) -> dict[str, Any] | None:
    raw = cell.value
    if raw is None and cached is None:
        return None
    data_type = getattr(cell, "data_type", None)
    is_formula = isinstance(raw, str) and raw.startswith("=")
    # A formula's value is ONLY its cached value. null means Excel stored none — the
    # formula string is never presented as a value (§16).
    value = cached if is_formula else (raw if raw is not None else cached)
    return {
        "cell": cell.coordinate,
        "row": cell.row,
        "column": cell.column,
        "data_type": data_type,
        "formula": raw if is_formula else None,
        "value": value,
    }


def _sheet_csv(rows: list[list[Any]]) -> str:
    buffer = io.StringIO(newline="")
    writer = csv.writer(buffer, lineterminator="\n")
    for row in rows:
        writer.writerow(["" if v is None else v for v in row])
    return buffer.getvalue()


def _rectangular_values(worksheet: Any, cached_ws: Any, max_row: int, max_col: int
                        ) -> tuple[list[list[Any]], bool]:
    """Value grid for CSV. Safe only when no cell holds a formula without a cached value."""
    grid: list[list[Any]] = []
    safe = True
    for r in range(1, max_row + 1):
        row: list[Any] = []
        for c in range(1, max_col + 1):
            raw = worksheet.cell(row=r, column=c).value
            cached = cached_ws.cell(row=r, column=c).value if cached_ws is not None else None
            if isinstance(raw, str) and raw.startswith("="):
                if cached is None:
                    safe = False
                    row.append(None)
                else:
                    row.append(cached)
            else:
                row.append(raw if raw is not None else cached)
        grid.append(row)
    return grid, safe


def _chart_metadata(worksheet: Any) -> list[dict[str, Any]]:
    out: list[dict[str, Any]] = []
    for chart in getattr(worksheet, "_charts", []) or []:
        try:
            title = None
            if getattr(chart, "title", None) is not None:
                title = str(chart.title)
            out.append({
                "chart_type": type(chart).__name__,
                "title": title,
                "series_count": len(getattr(chart, "series", []) or []),
            })
        except Exception:
            continue
    return out


def run(context: AdapterContext) -> ExtractionResult:
    fmt = context.detection.fmt.value
    result = ExtractionResult(document_id=context.document_id, format=fmt, adapter="xlsx")
    source = context.original_path

    try:
        import openpyxl
    except ImportError:
        result.fail("OPENPYXL_UNAVAILABLE")
        return result

    try:
        workbook = openpyxl.load_workbook(str(source), data_only=False, read_only=False)
    except Exception as exc:
        result.fail(f"XLSX_OPEN_FAILED:{type(exc).__name__}: {exc}")
        return result
    try:
        cached_workbook = openpyxl.load_workbook(str(source), data_only=True, read_only=False)
    except Exception:
        cached_workbook = None
        result.warn("XLSX_CACHED_VALUES_UNAVAILABLE")

    properties = workbook.properties
    workbook_meta = {
        "title": getattr(properties, "title", None),
        "creator": getattr(properties, "creator", None),
        "last_modified_by": getattr(properties, "lastModifiedBy", None),
        "created": getattr(properties, "created", None),
        "modified": getattr(properties, "modified", None),
    }
    named_ranges: list[dict[str, Any]] = []
    try:
        for name, defn in workbook.defined_names.items():
            named_ranges.append({"name": name, "value": getattr(defn, "value", None)})
    except Exception:
        result.warn("XLSX_NAMED_RANGES_UNAVAILABLE")

    builder = ElementBuilder()
    sheets_dir = context.bundle / SHEETS_DIRNAME
    sheet_records: list[dict[str, Any]] = []
    total_formulas = 0

    for index, sheet_name in enumerate(workbook.sheetnames, start=1):
        worksheet = workbook[sheet_name]
        cached_ws = cached_workbook[sheet_name] if cached_workbook is not None else None
        max_row = int(worksheet.max_row or 0)
        max_col = int(worksheet.max_column or 0)
        truncated = False
        if max_row * max_col > MAX_CELLS_PER_SHEET:
            truncated = True
            result.warn(f"XLSX_SHEET_TRUNCATED:{sheet_name}")
            max_row = min(max_row, max(1, MAX_CELLS_PER_SHEET // max(1, max_col)))

        cells: list[dict[str, Any]] = []
        for row in worksheet.iter_rows(min_row=1, max_row=max_row, max_col=max_col):
            for cell in row:
                cached = None
                if cached_ws is not None:
                    try:
                        cached = cached_ws.cell(row=cell.row, column=cell.column).value
                    except Exception:
                        cached = None
                record = _cell_record(cell, cached)
                if record is not None:
                    record["sheet"] = sheet_name
                    cells.append(record)
        formulas = [c for c in cells if c["formula"]]
        total_formulas += len(formulas)

        excel_tables: list[dict[str, Any]] = []
        try:
            for name, ref in (worksheet.tables or {}).items():
                excel_tables.append({"name": name, "ref": str(getattr(ref, "ref", ref))})
        except Exception:
            pass

        merged = [str(r) for r in (worksheet.merged_cells.ranges or [])]
        used_range = (f"A1:{worksheet.cell(row=max_row or 1, column=max_col or 1).coordinate}"
                      if max_row and max_col else None)

        grid, csv_safe = _rectangular_values(worksheet, cached_ws, max_row, max_col) \
            if max_row and max_col else ([], False)

        payload = {
            "document_id": context.document_id,
            "sheet": sheet_name,
            "sheet_index": index,
            "hidden": worksheet.sheet_state != "visible",
            "sheet_state": worksheet.sheet_state,
            "used_range": used_range,
            "max_row": max_row,
            "max_column": max_col,
            "truncated": truncated,
            "merged_cells": merged,
            "excel_tables": excel_tables,
            "charts": _chart_metadata(worksheet),
            "formula_count": len(formulas),
            "cells": cells,
        }
        sheets_dir.mkdir(parents=True, exist_ok=True)
        filename = safe_sheet_filename(sheet_name, index)
        json_path = sheets_dir / f"{filename}.json"
        json_path.write_text(json.dumps(payload, ensure_ascii=False, indent=2, default=str),
                             encoding="utf-8")
        csv_path: Path | None = None
        if grid and csv_safe and not merged:
            csv_path = sheets_dir / f"{filename}.csv"
            csv_path.write_text(_sheet_csv(grid), encoding="utf-8")
        elif grid and not csv_safe:
            result.warn(f"XLSX_CSV_SKIPPED_UNRESOLVED_FORMULAS:{sheet_name}")
        elif grid and merged:
            result.warn(f"XLSX_CSV_SKIPPED_MERGED_CELLS:{sheet_name}")

        record = base_asset(context.document_id, asset_id("SHEET", index), "SHEET", json_path)
        record.update({
            "sheet_name": sheet_name,
            "sheet_index": index,
            "hidden": payload["hidden"],
            "used_range": used_range,
            "rows": max_row,
            "columns": max_col,
            "merged_cells": merged,
            "excel_tables": excel_tables,
            "charts": payload["charts"],
            "formula_count": len(formulas),
            "cell_count": len(cells),
            "structured_path": relpath(json_path),
            "csv_path": relpath(csv_path) if csv_path else None,
            "snapshot_asset_id": None,
        })
        sheet_records.append(record)

        builder.add(HEADING, sheet_name, level=1, sheet_name=sheet_name,
                    cell_range=used_range, asset_id=record["asset_id"])
        preview = "\n".join(
            " | ".join("" if v is None else str(v) for v in row) for row in grid[:50]
        ).strip()
        if preview:
            builder.add(PARAGRAPH, preview, sheet_name=sheet_name, cell_range=used_range,
                        asset_id=record["asset_id"])

    # Optional, non-authoritative visual snapshots (§18).
    snapshots: list[dict[str, Any]] = []
    renderer = context.office_renderer
    if context.snapshots_enabled and renderer is not None:
        try:
            snapshots, snapshot_warnings = renderer.render_pages(
                source, context.bundle, document_id=context.document_id,
                scale=context.snapshot_scale, dirname=SHEET_SNAPSHOT_DIRNAME,
                prefix="sheet", kind="SHEET_SNAPSHOT", id_prefix="SHEETIMG")
            for warning in snapshot_warnings:
                result.warn(warning)
        except Exception as exc:  # never let a renderer defect destroy the extraction (§18)
            result.warn(f"SHEET_SNAPSHOT_FAILED_NONBLOCKING:{type(exc).__name__}")
    elif context.snapshots_enabled:
        result.warn(VISUAL_RENDERER_UNAVAILABLE)

    markdown_parts = [f"# {context.display_name}", ""]
    for record in sheet_records:
        markdown_parts.append(f"## {record['sheet_name']}"
                              + (" (gizli sayfa)" if record["hidden"] else ""))
        markdown_parts.append("")
        markdown_parts.append(
            f"- Kullanılan aralık: {record['used_range'] or 'yok'}"
        )
        markdown_parts.append(f"- Hücre sayısı: {record['cell_count']}")
        markdown_parts.append(f"- Formül sayısı: {record['formula_count']}")
        if record["excel_tables"]:
            markdown_parts.append(
                "- Excel tabloları: " + ", ".join(t["name"] for t in record["excel_tables"]))
        markdown_parts.append("")
    markdown = "\n".join(markdown_parts)

    document_json = {
        "document_id": context.document_id,
        "format": fmt,
        "adapter": "xlsx",
        "source_filename": context.original_filename,
        "structural_authority": "openpyxl_workbook_model",
        "workbook_properties": workbook_meta,
        "named_ranges": named_ranges,
        "sheet_names": list(workbook.sheetnames),
        "hidden_sheets": [r["sheet_name"] for r in sheet_records if r["hidden"]],
        "sheets": sheet_records,
        "elements": builder.elements,
    }
    text = "\n\n".join(e["text"] for e in builder.elements if e["text"])
    write_normalized(context.bundle, result, markdown=markdown,
                     document_json=document_json, text=text)

    if not sheet_records:
        result.fail("XLSX_NO_SHEETS_EXTRACTED")

    result.sheets = sheet_records
    result.pages = snapshots
    result.text_elements = builder.elements
    result.page_count = len(sheet_records)
    result.engine = {
        "native_parser": "openpyxl",
        "data_only_pass": cached_workbook is not None,
        "formula_count": total_formulas,
    }
    result.capabilities = {
        "text": bool(text.strip()),
        "tables": bool(sheet_records),
        "figures": False,
        "page_snapshots": bool(snapshots),
        "ocr": False,
        "vision": False,
        "formulas": total_formulas > 0,
        "charts": any(r["charts"] for r in sheet_records),
        "slides": False,
        "sheets": True,
    }
    return result
