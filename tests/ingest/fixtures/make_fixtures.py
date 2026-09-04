#!/usr/bin/env python3
"""Generate small deterministic synthetic fixtures for ingest tests (task §65).

Nothing here is derived from user/production documents. Run with the project venv:
    .venv/bin/python tests/ingest/fixtures/make_fixtures.py
"""

from __future__ import annotations

from pathlib import Path

HERE = Path(__file__).resolve().parent


def _scratch_png() -> Path:
    """A tiny synthetic raster reused as the embedded image in the DOCX/PPTX fixtures."""
    from PIL import Image, ImageDraw

    target = HERE / "_embedded.png"
    if not target.is_file():
        image = Image.new("RGB", (160, 120), "white")
        draw = ImageDraw.Draw(image)
        draw.rectangle([10, 10, 150, 110], outline="black", width=2)
        draw.text((24, 50), "ENKESIT", fill="black")
        image.save(str(target))
    return target


def make_docx(path: Path) -> None:
    from docx import Document

    d = Document()
    d.add_heading("NATM Destek Sistemi", level=1)
    d.add_paragraph("Bu sentetik test belgesidir. Püskürtme beton ve kaya bulonu.")
    d.add_heading("Tablo", level=2)
    t = d.add_table(rows=2, cols=2)
    t.cell(0, 0).text = "Eleman"
    t.cell(0, 1).text = "Kalınlık"
    t.cell(1, 0).text = "Püskürtme beton"
    t.cell(1, 1).text = "20 cm"
    d.add_heading("Şekil", level=2)
    d.add_paragraph("Şekil 1. Sentetik enkesit görseli.")
    d.add_picture(str(_scratch_png()))
    d.save(str(path))


def make_pptx(path: Path) -> None:
    from pptx import Presentation
    from pptx.util import Inches

    p = Presentation()
    s = p.slides.add_slide(p.slide_layouts[1])
    s.shapes.title.text = "Ovit Tuneli"
    s.placeholders[1].text = "Sentetik sunum. Slayt 1."
    s2 = p.slides.add_slide(p.slide_layouts[5])
    s2.shapes.title.text = "Maliyetler"
    box = s2.shapes.add_textbox(Inches(1), Inches(2), Inches(4), Inches(1))
    box.text_frame.text = "Toplam: 100"
    s2.shapes.add_picture(str(_scratch_png()), Inches(1), Inches(3), Inches(2))
    s2.notes_slide.notes_text_frame.text = "Konusmaci notu: maliyet dagilimi."
    p.save(str(path))


def make_xlsx(path: Path) -> None:
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = "Maliyetler"
    ws["A1"] = "Kalem"
    ws["B1"] = "Tutar"
    ws["A2"] = "Kazi"
    ws["B2"] = 40
    ws["A3"] = "Destek"
    ws["B3"] = 60
    ws["A4"] = "Toplam"
    ws["B4"] = "=SUM(B2:B3)"
    ws["D1"] = "Not"
    ws.merge_cells("D2:E2")
    ws["D2"] = "birlesik hucre"
    ozet = wb.create_sheet("Ozet")          # values only -> a CSV is safe to emit (§17)
    for row, values in enumerate([["Yil", "Tutar"], [2024, 100], [2025, 120]], start=1):
        for col, value in enumerate(values, start=1):
            ozet.cell(row=row, column=col, value=value)
    hidden = wb.create_sheet("Gizli")
    hidden["A1"] = "gizli sayfa"
    hidden.sheet_state = "hidden"
    wb.save(str(path))


def make_png_with_text(path: Path) -> None:
    from PIL import Image, ImageDraw

    img = Image.new("RGB", (600, 200), "white")
    draw = ImageDraw.Draw(img)
    draw.text((20, 40), "Puskurtme Beton 20 cm", fill="black")
    draw.text((20, 90), "Kaya Bulonu", fill="black")
    draw.text((20, 140), "Celik Hasir", fill="black")
    img.save(str(path))


def make_jpg_photo(path: Path) -> None:
    from PIL import Image

    # a gradient — no text, so OCR must NOT hallucinate (task §70)
    img = Image.new("RGB", (320, 240))
    px = img.load()
    for y in range(240):
        for x in range(320):
            px[x, y] = (x % 256, y % 256, (x + y) % 256)
    img.save(str(path), quality=70)


def make_pdf(path: Path) -> None:
    from PIL import Image, ImageDraw

    pages = []
    for n in (1, 2):
        img = Image.new("RGB", (612, 792), "white")
        d = ImageDraw.Draw(img)
        d.text((72, 72), f"Tunel Bakim Raporu - Sayfa {n}", fill="black")
        d.text((72, 120), "Sentetik PDF icerigi. NATM ve TBM.", fill="black")
        pages.append(img)
    pages[0].save(str(path), save_all=True, append_images=pages[1:])


def make_pdf_rich(path: Path) -> None:
    """A hand-written PDF with a real text layer, a ruled table and an embedded raster figure.

    Authored byte-by-byte because the venv has no PDF-authoring library; everything in it is
    synthetic. `sample.pdf` stays image-only on purpose so the OCR fallback is exercised too.
    """
    import zlib

    def _esc(text: str) -> str:
        return text.replace("\\", r"\\").replace("(", r"\(").replace(")", r"\)")

    width, height = 612, 792

    # --- page 1: headings, paragraphs, a ruled table -------------------------
    lines1 = []
    def txt(x, y, size, s, font="F1"):
        lines1.append(f"BT /{font} {size} Tf 1 0 0 1 {x} {y} Tm ({_esc(s)}) Tj ET")

    txt(72, 720, 18, "Tunel Bakim Maliyet Raporu", "F2")
    txt(72, 690, 11, "1. Giris")
    txt(72, 672, 10, "Bu sentetik test belgesi tunel bakim maliyetlerini ozetler.")
    txt(72, 656, 10, "NATM ve TBM yontemleri karsilastirilmistir.")
    txt(72, 620, 11, "2. Maliyet Tablosu")

    # table grid: 3 cols x 4 rows
    x0, y0, cw, rh, cols, rows = 72, 480, 140, 24, 3, 4
    for r in range(rows + 1):
        y = y0 + r * rh
        lines1.append(f"0.5 w 0 G {x0} {y} m {x0 + cols * cw} {y} l S")
    for c in range(cols + 1):
        x = x0 + c * cw
        lines1.append(f"0.5 w 0 G {x} {y0} m {x} {y0 + rows * rh} l S")
    cells = [
        ["Kalem", "Birim", "Tutar"],
        ["Kazi", "m3", "40"],
        ["Destek", "m2", "60"],
        ["Toplam", "-", "100"],
    ]
    for r, row in enumerate(cells):
        y = y0 + (rows - r - 1) * rh + 8
        for c, value in enumerate(row):
            txt(x0 + c * cw + 6, y, 10, value, "F2" if r == 0 else "F1")
    txt(72, 440, 9, "Tablo 1. Birim maliyet ozeti.")

    content1 = "\n".join(lines1).encode("latin-1")

    # --- page 2: an embedded raster figure + caption -------------------------
    iw, ih = 120, 90
    pixels = bytearray()
    for y in range(ih):
        for x in range(iw):
            pixels += bytes((40 + (x * 2) % 200, 60, 200 - (y * 2) % 180))
    image_stream = zlib.compress(bytes(pixels))

    lines2 = []
    lines2.append(f"BT /F2 14 Tf 1 0 0 1 72 720 Tm ({_esc('3. Kesit Gorseli')}) Tj ET")
    lines2.append("q 240 0 0 180 72 500 cm /Im1 Do Q")
    lines2.append(f"BT /F1 9 Tf 1 0 0 1 72 480 Tm ({_esc('Sekil 1. Tunel enkesiti semasi.')}) Tj ET")
    lines2.append(f"BT /F1 10 Tf 1 0 0 1 72 440 Tm ({_esc('Gorsel yerel olarak uretilmistir; gercek proje verisi degildir.')}) Tj ET")
    content2 = "\n".join(lines2).encode("latin-1")

    objects: list[bytes] = []

    def add(body: bytes) -> int:
        objects.append(body)
        return len(objects)

    catalog = add(b"")            # 1
    pages = add(b"")              # 2
    page1 = add(b"")              # 3
    content1_obj = add(b"")       # 4
    page2 = add(b"")              # 5
    content2_obj = add(b"")       # 6
    font1 = add(b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>")          # 7
    font2 = add(b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica-Bold >>")     # 8
    image = add(b"")              # 9

    objects[catalog - 1] = f"<< /Type /Catalog /Pages {pages} 0 R >>".encode()
    objects[pages - 1] = (
        f"<< /Type /Pages /Kids [{page1} 0 R {page2} 0 R] /Count 2 >>".encode())
    resources = (f"<< /Font << /F1 {font1} 0 R /F2 {font2} 0 R >> "
                 f"/XObject << /Im1 {image} 0 R >> >>")
    objects[page1 - 1] = (
        f"<< /Type /Page /Parent {pages} 0 R /MediaBox [0 0 {width} {height}] "
        f"/Resources {resources} /Contents {content1_obj} 0 R >>").encode()
    objects[page2 - 1] = (
        f"<< /Type /Page /Parent {pages} 0 R /MediaBox [0 0 {width} {height}] "
        f"/Resources {resources} /Contents {content2_obj} 0 R >>").encode()
    objects[content1_obj - 1] = (
        f"<< /Length {len(content1)} >>\nstream\n".encode() + content1 + b"\nendstream")
    objects[content2_obj - 1] = (
        f"<< /Length {len(content2)} >>\nstream\n".encode() + content2 + b"\nendstream")
    objects[image - 1] = (
        f"<< /Type /XObject /Subtype /Image /Name /Im1 /Width {iw} /Height {ih} "
        f"/ColorSpace /DeviceRGB /BitsPerComponent 8 /Filter /FlateDecode "
        f"/Length {len(image_stream)} >>\nstream\n".encode() + image_stream + b"\nendstream")

    out = bytearray(b"%PDF-1.4\n%\xe2\xe3\xcf\xd3\n")
    offsets = [0]
    for number, body in enumerate(objects, start=1):
        offsets.append(len(out))
        out += f"{number} 0 obj\n".encode() + body + b"\nendobj\n"
    xref_pos = len(out)
    out += f"xref\n0 {len(objects) + 1}\n".encode()
    out += b"0000000000 65535 f \n"
    for offset in offsets[1:]:
        out += f"{offset:010d} 00000 n \n".encode()
    out += (f"trailer\n<< /Size {len(objects) + 1} /Root {catalog} 0 R >>\n"
            f"startxref\n{xref_pos}\n%%EOF\n").encode()
    Path(path).write_bytes(bytes(out))


def make_txt(path: Path) -> None:
    path.write_text("Tunel tanimi: yeraltinda insa edilen gecis yapisi.\n", encoding="utf-8")


BUILDERS = {
    "sample.docx": make_docx,
    "sample.pptx": make_pptx,
    "sample.xlsx": make_xlsx,
    "sample_text.png": make_png_with_text,
    "sample_photo.jpg": make_jpg_photo,
    "sample.pdf": make_pdf,
    "sample_rich.pdf": make_pdf_rich,
    "sample_text.txt": make_txt,
}


def build_all(dest: Path | None = None) -> list[Path]:
    dest = dest or HERE
    dest.mkdir(parents=True, exist_ok=True)
    out = []
    for name, fn in BUILDERS.items():
        target = dest / name
        fn(target)
        out.append(target)
    return out


if __name__ == "__main__":
    for p in build_all():
        print(p, p.stat().st_size)
